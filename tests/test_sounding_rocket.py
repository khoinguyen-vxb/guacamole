"""Launcher checks; backend engineering and paid providers are never invoked."""

import importlib.util
import io
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from runtime_support import TestReasoning

from examples.sounding_rocket import kookaburra
from guacamole import (
    AgentSpec,
    AgentTurn,
    Budget,
    ContextSpec,
    Project,
    ProjectRequest,
    ToolRegistry,
)
from guacamole.contracts import Rationale, ToolCall, Wait


class SoundingRocketTests(unittest.IsolatedAsyncioTestCase):
    async def test_resume_refreshes_limits_without_replacing_saved_context(self):
        class Idle(TestReasoning):
            async def respond(self, packet):
                return AgentTurn(
                    rationale=Rationale(objective="Wait", explanation="Fixture"),
                    action=Wait(reason="Fixture paused"),
                )

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with Project(root, reasoning=Idle()) as project:
                result = await project.run(
                    ProjectRequest(
                        description="Preserve this brief",
                        context=ContextSpec(
                            max_input_tokens=128_000,
                            instructions="Preserve instructions",
                        ),
                    )
                )
                before = project.store.model(
                    "agent", project.store.list("agent")[0].ref.id, AgentSpec
                )
            budget = Budget(max_steps=300, max_tool_calls=250, max_seconds=9000)
            with (
                Project(root, reasoning=Idle()) as project,
                patch.object(
                    kookaburra,
                    "CONTEXT",
                    ContextSpec(
                        max_input_tokens=12_800_000, reserved_output_tokens=20_000
                    ),
                ),
                patch.object(kookaburra, "BUDGET", budget),
                redirect_stdout(io.StringIO()) as output,
            ):
                context = kookaburra.configure_run(project, result.run_id)
                self.assertEqual(context.max_input_tokens, 480_000)
                self.assertIn("480000", output.getvalue())
                saved = project.store.model("agent", before.id, AgentSpec)
                self.assertEqual(
                    saved.context,
                    before.context.model_copy(
                        update={
                            "max_input_tokens": 480_000,
                            "reserved_output_tokens": 20_000,
                        }
                    ),
                )
                self.assertEqual(saved.budget, budget)
                self.assertEqual(saved.decision_id, before.decision_id)
                await project.run(resume=result.run_id)
                request = project.store.get("run", result.run_id).data["request"]
                self.assertEqual(request["description"], "Preserve this brief")
                self.assertEqual(request["context"]["max_input_tokens"], 480_000)
                self.assertEqual(request["budget"], budget.model_dump(mode="json"))
                events = len(project.events())
                kookaburra.configure_run(project, result.run_id)
                self.assertEqual(len(project.events()), events)

    async def test_document_reads_before_planning_preserve_grants_and_write_gate(self):
        class Reader(TestReasoning):
            async def respond(self, packet):
                return AgentTurn(
                    rationale=Rationale(
                        objective="Inspect a source", explanation="Test fixture only"
                    ),
                    action=next(actions),
                )

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with ZipFile(root / "previous.ork", "w") as archive:
                archive.writestr("rocket.ork", "<rocket>Historical design</rocket>")
            with Project(
                root / "state", sandbox=root / "outputs", reasoning=Reader()
            ) as project:
                kookaburra.register_tools(project)
                source = project.ingest(root / "previous.ork")
                reader = project.tools.get("read_document@1")
                arguments = {"source": source.model_dump(mode="json")}
                actions = iter(
                    (
                        ToolCall(
                            tool=reader.definition.grant,
                            arguments=arguments,
                            inputs=(source,),
                        ),
                        ToolCall(
                            tool=project.tools.get("write_text@1").definition.grant,
                            arguments={
                                "path": "unplanned.txt",
                                "text": "must not be written",
                            },
                        ),
                        Wait(reason="Fixture finished"),
                    )
                )
                await project.run(
                    ProjectRequest(
                        description="Source inspection fixture",
                        context=ContextSpec(max_input_tokens=200_000),
                    )
                )
                jobs = [r for r in project.requests() if r.kind == "tool"]
                self.assertEqual(
                    [(r.target, r.status, r.decision_id) for r in jobs],
                    [("read_document@1", "completed", None)],
                )
                chief = project.store.list("agent")[0]
                excerpt = project.store.resolve(
                    project.extract(jobs[0].id, actor=chief.ref.id)
                ).data
                self.assertIn("Historical design", excerpt["text"])
                self.assertEqual(excerpt["source"], source.model_dump(mode="json"))
                self.assertFalse(project.store.list("decision"))
                self.assertFalse((project.sandbox / "unplanned.txt").exists())
                with self.assertRaisesRegex(PermissionError, "not granted"):
                    project.executor.submit(
                        reader, arguments, actor=chief.ref.id, grants=()
                    )
                for name in ("list_files", "read_text", "read_document"):
                    self.assertFalse(
                        project.tools.get(f"{name}@1").definition.requires_plan
                    )
                for name in ("write_text", "execute", "verify_command"):
                    with (
                        self.subTest(tool=name),
                        self.assertRaisesRegex(PermissionError, "validated work plan"),
                    ):
                        project.executor.submit(
                            project.tools.get(f"{name}@1"), {}, actor=chief.ref.id
                        )

    async def test_registered_tools_execute_capture_verify_and_reject_escapes(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            repo = root / "tooling" / "fixture"
            repo.mkdir(parents=True)
            (repo / "engine.py").write_text("value = 7\n")
            (repo / "README.md").write_text(
                "Fixture backend, not engineering evidence."
            )
            documents = root / "technical_documents"
            documents.mkdir()
            (documents / "brief.txt").write_text("Historical input")
            with (
                patch.multiple(
                    kookaburra,
                    ROOT=root,
                    TOOLING=repo.parent,
                    DOCUMENTS=documents,
                    EXTRA_TOOLS=(),
                ),
                Project(
                    root / "state", sandbox=root / "outputs", tools=ToolRegistry()
                ) as project,
            ):
                kookaburra.register_tools(project)
                request = kookaburra.project_request(
                    project, ContextSpec(max_input_tokens=128_000)
                )
                self.assertIn(documents / "brief.txt", request.documents)
                self.assertIn(repo / "README.md", request.documents)
                self.assertIn("30000", request.description.replace(",", ""))
                self.assertEqual(request.context.max_input_tokens, 128_000)
                self.assertEqual(len(project.tools.definitions()), 6)
                write = project.tools.get("write_text@1")
                await project.call(
                    write,
                    {
                        "path": "work/calculate.py",
                        "text": "import engine\nfrom pathlib import Path\nPath('value.txt').write_text(str(engine.value))\n",
                    },
                )
                execution = await project.call(
                    project.tools.get("execute@1"),
                    {
                        "command": ["python", "calculate.py"],
                        "repository": "fixture",
                        "outputs": ["work/value.txt", "work/missing.txt"],
                    },
                )
                self.assertEqual(execution.returncode, 0)
                self.assertFalse(execution.timed_out)
                self.assertEqual(Path(execution.files[0].path).read_text(), "7")
                self.assertEqual(len(execution.missing_outputs), 1)
                self.assertTrue(project.store.list("job_file"))
                verification = await project.call(
                    project.tools.get("verify_command@1"),
                    {
                        "command": [
                            "python",
                            "-c",
                            "import json; from pathlib import Path; passed = Path('value.txt').read_text() == '7'; print(json.dumps({'outcome': 'passed' if passed else 'failed', 'explanation': 'Compared actual file', 'checks': {'value': passed}}))",
                        ]
                    },
                )
                self.assertEqual(verification.outcome, "passed")
                timed_out = await project.call(
                    project.tools.get("execute@1"),
                    {
                        "command": ["python", "-c", "import time; time.sleep(30)"],
                        "timeout_s": 0.05,
                    },
                )
                self.assertTrue(timed_out.timed_out)
                with self.assertRaises(ValueError):
                    await project.call(write, {"path": "../escape.txt", "text": "no"})
                with self.assertRaises(ValueError):
                    await project.call(
                        project.tools.get("read_text@1"),
                        {"path": str(root / "state/project.sqlite3")},
                    )
                with self.assertRaises(ValueError):
                    project.tools.get("execute@1").arguments(
                        {"command": "python calculate.py"}
                    )

    @unittest.skipUnless(
        all(importlib.util.find_spec(name) for name in ("pypdf", "openpyxl", "pptx")),
        "Install the sounding-rocket extra to check its document readers",
    )
    async def test_document_readers_preserve_locators_and_provenance(self):
        from openpyxl import Workbook
        from pptx import Presentation
        from pypdf import PdfWriter

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Payload"
            worksheet.append(["mass_kg", 2.5])
            workbook.save(root / "requirements.xlsx")
            slides = Presentation()
            slide = slides.slides.add_slide(slides.slide_layouts[1])
            slide.shapes.title.text = "Payload interface"
            slide.notes_slide.notes_text_frame.text = "Check the connector"
            slides.save(root / "review.pptx")
            pdf = PdfWriter()
            pdf.add_blank_page(width=100, height=100)
            pdf.write(root / "drawing.pdf")
            with ZipFile(root / "previous.ork", "w") as archive:
                archive.writestr(
                    "rocket.ork", "<rocket><name>Previous design</name></rocket>"
                )
            with Project(root / "state", sandbox=root / "outputs") as project:
                kookaburra.register_tools(project)
                reader = project.tools.get("read_document@1")
                for filename, expected, arguments in (
                    ("requirements.xlsx", '"Payload"', {}),
                    ("requirements.xlsx", "B1=2.5", {"sheet": "Payload"}),
                    ("review.pptx", "Check the connector", {}),
                    ("drawing.pdf", "[No extractable text]", {}),
                    ("previous.ork", "Previous design", {}),
                ):
                    source = project.ingest(root / filename)
                    job = project.submit(
                        reader,
                        {"source": source.model_dump(mode="json"), **arguments},
                        inputs=(source,),
                    )
                    excerpt = await job.collect()
                    self.assertIn(expected, excerpt.text)
                    self.assertIn(filename, excerpt.locator)
                    ref = project.extract(job.id, actor="test")
                    self.assertEqual(
                        project.store.resolve(ref).data["source"],
                        source.model_dump(mode="json"),
                    )
                source = project.ingest(root / "drawing.pdf")
                with self.assertRaises(ValueError):
                    await project.call(
                        reader, {"source": source.model_dump(mode="json"), "first": 2}
                    )

    async def test_launch_uses_codex_without_api_key(self):
        with (
            patch.dict(kookaburra.os.environ, {}, clear=True),
            patch.object(kookaburra.shutil, "which", return_value="/usr/bin/codex"),
            patch.object(kookaburra.importlib.util, "find_spec", return_value=True),
            patch.object(
                kookaburra, "Project", side_effect=RuntimeError("startup reached")
            ) as project,
            self.assertRaisesRegex(RuntimeError, "startup reached"),
        ):
            await kookaburra.main()
        self.assertIsInstance(
            project.call_args.kwargs["reasoning"], kookaburra.CodexReasoning
        )

    async def test_launch_requires_readers_without_creating_workspace(self):
        with tempfile.TemporaryDirectory() as directory:
            workspace = Path(directory) / "workspace"
            with (
                patch.object(kookaburra, "WORKSPACE", workspace),
                patch.object(kookaburra.shutil, "which", return_value="/usr/bin/codex"),
                patch.object(kookaburra.importlib.util, "find_spec", return_value=None),
            ):
                with self.assertRaisesRegex(
                    SystemExit, "uv sync --extra sounding-rocket"
                ):
                    await kookaburra.main()
                self.assertFalse(workspace.exists())

    async def test_launch_requires_codex_without_creating_workspace(self):
        with tempfile.TemporaryDirectory() as directory:
            workspace = Path(directory) / "workspace"
            with (
                patch.object(
                    kookaburra, "ROOT", Path(directory) / "examples/sounding_rocket"
                ),
                patch.object(kookaburra, "WORKSPACE", workspace),
                patch.dict(kookaburra.os.environ, {}, clear=True),
                patch.object(kookaburra.shutil, "which", return_value=None),
            ):
                with self.assertRaisesRegex(SystemExit, "codex login"):
                    await kookaburra.main()
                self.assertFalse(workspace.exists())


if __name__ == "__main__":
    unittest.main()
